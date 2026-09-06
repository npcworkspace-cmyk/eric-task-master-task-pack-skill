#!/usr/bin/env python3
"""Offline audit/export for collector artifacts. No network access or real-data fixtures.

python audit_export.py --source RUN_DIR --output DELIVERY_DIR [--baseline OLD.json] [--xlsx]
Exit 0: no data-validity errors; 2: invalid data or fatal input/output error.
Exit 0 does NOT mean complete coverage or completed human review: read audit.json.
"""
from __future__ import annotations

import argparse
import copy
import csv
import importlib.util
import json
import math
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit, urlunsplit

BUCKETS = ('posts', 'outside_range_posts', 'undated_posts')
METRICS = ('shares', 'reactions', 'comments')
COLUMNS = [
    ('record_bucket','原始分组'), ('record_index','分组内序号'), ('post_id','贴文ID'),
    ('url','贴文链接'), ('published_at','发布时间（ISO含时区）'), ('published_at_raw','原始时间'),
    ('body','正文'), ('shared_body','被转发原帖正文'), ('classification','正文类型'),
    ('shares','分享数'), ('reactions','点赞心情数'), ('comments','评论数'),
    ('shares_raw','分享原始标签'), ('reactions_raw','心情原始标签'), ('comments_raw','评论原始标签'),
    ('count_observed_at.shares','分享数观测时间'), ('count_observed_at.reactions','心情数观测时间'), ('count_observed_at.comments','评论数观测时间'),
    ('date_status','日期状态'), ('row_issues','待核验项'), ('collected_at','采集时间（ISO含时区）'),
    ('journal_status','分页记录保真状态'),
]
COMPACT = re.compile(r'(?<![\w.])\d+(?:[.,]\d+)?\s*(?:[kKmMbB]|万|萬|千|亿|億)(?![A-Za-z])')
INTEGER = re.compile(r'(?<![\d.])(?:\d{1,3}(?:[,\u00a0 ]\d{3})+|\d+)(?![\d.])')
ILLEGAL_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def iso(value):
    if not isinstance(value, str) or not re.match(r'^\d{4}-\d{2}-\d{2}T', value):
        return None
    try:
        parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
        return parsed if parsed.tzinfo is not None else None
    except ValueError:
        return None


def integer(value):
    return isinstance(value, int) and not isinstance(value, bool)


def scalar(value):
    return value is None or isinstance(value, (str, int, float, bool))


def clean_retry_observation(record, kind):
    """Match only the collector's bounded retry evidence, never arbitrary failures."""
    audit = record.get('audit')
    posts = record.get('posts')
    if (not integer(record.get('page')) or record['page'] < 1
            or not isinstance(record.get('request_cursor'), str) or not record['request_cursor']
            or not isinstance(audit, dict) or not isinstance(posts, list) or record.get('errors')
            or any(not isinstance(p, dict) or not isinstance(p.get('post_id'), str)
                   or not p['post_id'].isdigit() for p in posts)):
        return False
    ids = [p['post_id'] for p in posts]
    if (len(set(ids)) != len(ids) or audit.get('kind') != 'page'
            or not integer(audit.get('page')) or audit['page'] != record['page']
            or audit.get('status') != record.get('status')
            or audit.get('stream_final') is not record.get('stream_final')
            or audit.get('has_next') is not record.get('has_next')
            or not integer(audit.get('http_status')) or audit['http_status'] != 200
            or not integer(audit.get('frame_count')) or audit['frame_count'] < 1
            or not integer(audit.get('malformed_frames')) or audit['malformed_frames'] != 0
            or not integer(audit.get('rejected_stories')) or audit['rejected_stories'] != 0 or audit.get('errors') != []
            or not integer(audit.get('count')) or audit['count'] != len(posts) or audit.get('ids') != ids
            or audit.get('dates') != [p.get('published_at') for p in posts]):
        return False
    if kind == 'incomplete':
        return (record.get('status') == 'pagination_incomplete' and record.get('stream_final') is False
                and record.get('has_next') is None and record.get('next_cursor') is None and 1 <= len(posts) <= 3)
    if kind == 'terminal':
        return (record.get('status') == 'ok' and record.get('stream_final') is True
                and record.get('has_next') is False and record.get('next_cursor') is None and not posts)
    if kind == 'advancing':
        return (record.get('status') == 'ok' and record.get('stream_final') is True
                and record.get('has_next') is True and len(posts) == 3
                and isinstance(record.get('next_cursor'), str) and bool(record['next_cursor'])
                and record['next_cursor'] != record['request_cursor'])
    return False


def table_text(value):
    """Only table exports receive text prefixes. JSON source text remains unchanged."""
    if isinstance(value, str) and (value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n'))):
        return "'" + value
    return value


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def group_token(url):
    if not isinstance(url, str):
        return None
    try:
        parts = urlsplit(url)
        match = re.match(r'^/groups/([^/]+)', parts.path)
        host = parts.hostname or ''
        if parts.scheme != 'https' or parts.username or parts.password or not (host == 'facebook.com' or host.endswith('.facebook.com')):
            return None
        return unquote(match.group(1)) if match else None
    except ValueError:
        return None


class Auditor:
    def __init__(self, source, baseline=None):
        self.source = Path(source).resolve()
        self.baseline = Path(baseline) if baseline else None
        self.issues = []
        self.data = read_json(self.source / 'posts.json')
        if not isinstance(self.data, dict):
            raise ValueError('POSTS_ROOT_NOT_OBJECT')
        self.meta = self.data.get('metadata') or {}
        if not isinstance(self.meta, dict):
            self.meta = {}
            self.issue('validity', 'error', 'METADATA_NOT_OBJECT')
        self.start, self.end = iso(self.meta.get('start')), iso(self.meta.get('end'))
        self.group_id = str(self.meta['group_id']) if scalar(self.meta.get('group_id')) and self.meta.get('group_id') is not None else None
        self.group_aliases = {self.group_id, group_token(self.meta.get('group_url'))} - {None, ''}
        if not self.group_id or not self.group_id.isdigit() or not self.start or not self.end or self.start > self.end:
            self.issue('validity', 'error', 'INVALID_SCOPE')
        if not group_token(self.meta.get('group_url')):
            self.issue('validity', 'error', 'INVALID_GROUP_URL')

    def issue(self, area, level, code, **location):
        self.issues.append({'area': area, 'level': level, 'code': code, **location})

    def same_scope(self, value):
        basic = (isinstance(value, dict) and str(value.get('group_id')) == self.group_id
                and iso(value.get('start')) == self.start and iso(value.get('end')) == self.end
                and self.start is not None and self.end is not None)
        if not basic:
            return False
        if value.get('sorting_setting', 'CHRONOLOGICAL') != 'CHRONOLOGICAL' or value.get('page_size', 3) != 3:
            return False
        if value.get('required_boundary_pages', 5) != self.meta.get('required_boundary_pages', 5):
            return False
        if value.get('schema_version') == 3 or 'operation' in value:
            return (value.get('group_url') == self.meta.get('group_url')
                    and value.get('operation') == self.meta.get('operation', 'GroupsCometFeedRegularStoriesPaginationQuery'))
        return True

    def text_field(self, row, name, loc):
        value = row.get(name)
        if value is not None and not isinstance(value, str):
            self.issue('validity', 'error', 'TEXT_FIELD_NOT_TEXT', field=name, **loc)
            return None
        return value

    def metric(self, value, label, name, loc):
        # A precise numeric collector field takes precedence over an abbreviated display label.
        if integer(value):
            if value >= 0:
                return value, 'exact'
            self.issue('validity', 'error', 'INVALID_COUNT', field=name, **loc)
            return value, 'invalid'
        if isinstance(value, float):
            if math.isfinite(value) and value >= 0 and float(value).is_integer():
                return int(value), 'exact'
            self.issue('validity', 'error', 'INVALID_COUNT', field=name, **loc)
            return value if math.isfinite(value) else None, 'invalid'
        if value is not None and not isinstance(value, str):
            self.issue('validity', 'error', 'INVALID_COUNT_TYPE', field=name, **loc)
            return None, 'invalid'
        text = value if isinstance(value, str) and value.strip() else label
        if not text:
            self.issue('review', 'warning', 'COUNT_NOT_AVAILABLE', field=name, **loc)
            return None, 'missing'
        short = COMPACT.search(text)
        if short:
            self.issue('review', 'warning', 'APPROXIMATE_COUNT', field=name, **loc)
            return '约 ' + short.group(0), 'approximate'
        matches = INTEGER.findall(text)
        if len(matches) == 1 and not re.search(r'-\s*\d|\d\.\d', text):
            return int(re.sub(r'[,\u00a0 ]', '', matches[0])), 'exact_label'
        self.issue('validity', 'error', 'UNREADABLE_COUNT', field=name, **loc)
        return value, 'invalid'

    def normalize(self):
        normalized = {}
        seen = {}
        for bucket in BUCKETS:
            values = self.data.get(bucket, [])
            if values is None:
                self.issue('validity', 'error', 'BUCKET_IS_NULL', bucket=bucket)
                values = []
            if not isinstance(values, list):
                self.issue('validity', 'error', 'BUCKET_NOT_ARRAY', bucket=bucket)
                values = []
            normalized[bucket] = []
            for index, raw in enumerate(values, 1):
                start_issue = len(self.issues)
                loc = {'bucket': bucket, 'index': index}
                if not isinstance(raw, dict):
                    self.issue('validity', 'error', 'POST_NOT_OBJECT', **loc)
                    raw = {}
                pid = raw.get('post_id')
                if integer(pid) and pid >= 0:
                    pid = str(pid)
                    self.issue('review', 'warning', 'NUMERIC_ID_COERCED_TO_TEXT', **loc)
                if not isinstance(pid, str) or not re.fullmatch(r'\d+', pid):
                    self.issue('validity', 'error', 'INVALID_POST_ID', **loc)
                    pid = pid if isinstance(pid, str) else None
                elif pid in seen:
                    self.issue('validity', 'error', 'DUPLICATE_POST_ID', post_id=pid, **loc)
                else:
                    seen[pid] = loc
                item = {'record_bucket': bucket, 'record_index': index, 'post_id': pid}
                url = self.text_field(raw, 'url', loc)
                safe_url = None
                try:
                    parts = urlsplit(url or '')
                    match = re.fullmatch(r'/groups/([^/]+)/(?:posts|permalink)/(\d+)/?', parts.path)
                    host = parts.hostname or ''
                    valid = (parts.scheme == 'https' and (host == 'facebook.com' or host.endswith('.facebook.com'))
                             and not parts.username and not parts.password and match is not None
                             and unquote(match.group(1)) in self.group_aliases and match.group(2) == pid)
                    if parts.scheme == 'https' and host and not parts.username and not parts.password:
                        safe_url = urlunsplit((parts.scheme, parts.netloc, parts.path, '', ''))
                except ValueError:
                    valid = False
                if not valid:
                    self.issue('validity', 'error', 'URL_ID_OR_GROUP_MISMATCH', **loc)
                item['url'] = safe_url
                for field in ('body', 'shared_body', 'published_at', 'collected_at', 'shares_raw', 'reactions_raw', 'comments_raw'):
                    item[field] = self.text_field(raw, field, loc)
                date = iso(item['published_at'])
                if not date:
                    self.issue('review', 'warning', 'UNKNOWN_PUBLICATION_DATE', **loc)
                    item['date_status'] = 'unknown'
                elif self.start and self.end:
                    item['date_status'] = 'in_range' if self.start <= date <= self.end else 'outside_range'
                    if bucket == 'posts' and item['date_status'] != 'in_range':
                        self.issue('validity', 'error', 'TARGET_POST_OUTSIDE_SCOPE', **loc)
                    if bucket == 'outside_range_posts' and item['date_status'] == 'in_range':
                        self.issue('validity', 'error', 'OUTSIDE_BUCKET_CONTAINS_TARGET_POST', **loc)
                    if bucket == 'undated_posts':
                        self.issue('review', 'warning', 'DATED_POST_IN_UNDATED_BUCKET', **loc)
                else:
                    item['date_status'] = 'scope_unknown'
                raw_date = raw.get('published_at_raw')
                item['published_at_raw'] = raw_date if scalar(raw_date) and not isinstance(raw_date, bool) else None
                if isinstance(raw_date, float) and not math.isfinite(raw_date):
                    item['published_at_raw'] = None
                if raw_date is not None:
                    if isinstance(raw_date, (int, float)) and not isinstance(raw_date, bool):
                        try:
                            epoch = datetime.fromtimestamp(raw_date, timezone.utc)
                            if date and abs((date - epoch).total_seconds()) > 0.001:
                                self.issue('validity', 'error', 'EPOCH_ISO_MISMATCH', **loc)
                        except (ValueError, OSError, OverflowError):
                            self.issue('validity', 'error', 'INVALID_EPOCH', **loc)
                    elif not isinstance(raw_date, str):
                        self.issue('validity', 'error', 'INVALID_RAW_TIME_TYPE', **loc)
                    elif re.fullmatch(r'\d+(?:\.\d+)?', raw_date):
                        if date and abs(date.timestamp() - float(raw_date)) > 0.001:
                            self.issue('validity', 'error', 'EPOCH_ISO_MISMATCH', **loc)
                if not iso(item['collected_at']):
                    self.issue('review', 'warning', 'UNKNOWN_COLLECTION_TIME', **loc)
                item['metric_status'] = {}
                for name in METRICS:
                    item[name], item['metric_status'][name] = self.metric(raw.get(name), item[name + '_raw'], name, loc)
                observed = raw.get('count_observed_at')
                if observed is not None and not isinstance(observed, dict):
                    self.issue('validity', 'error', 'INVALID_COUNT_OBSERVATION_OBJECT', **loc)
                    observed = {}
                item['count_observed_at'] = {}
                for name in METRICS:
                    value = (observed or {}).get(name)
                    if value is not None and not iso(value):
                        self.issue('validity','error','INVALID_COUNT_OBSERVATION_TIME',field=name,**loc)
                        value = None
                    if value is None and item[name] is not None:
                        self.issue('review','warning','COUNT_OBSERVATION_TIME_UNAVAILABLE',field=name,**loc)
                    item['count_observed_at'][name] = value
                field_times = raw.get('field_observed_at')
                if field_times is not None and not isinstance(field_times,dict):
                    self.issue('validity','error','INVALID_FIELD_OBSERVATION_OBJECT',**loc)
                    field_times = {}
                item['field_observed_at'] = {}
                for name in ('body','url','publication_time'):
                    value = (field_times or {}).get(name)
                    if value is not None and not iso(value):
                        self.issue('validity','error','INVALID_FIELD_OBSERVATION_TIME',field=name,**loc)
                        value = None
                    item['field_observed_at'][name] = value
                item['retained_fields'] = []
                retained = raw.get('retained_fields')
                if retained is not None and not isinstance(retained,list):
                    self.issue('validity','error','INVALID_RETAINED_FIELDS',**loc)
                    retained=[]
                for previous_field in retained or []:
                    if (not isinstance(previous_field,dict) or previous_field.get('field') not in ('body','url','published_at')
                            or previous_field.get('reason') != 'unavailable_in_new_observation' or not iso(previous_field.get('from_collected_at'))):
                        self.issue('validity','error','INVALID_FIELD_RETENTION_RECORD',**loc)
                        continue
                    item['retained_fields'].append({key:previous_field[key] for key in ('field','from_collected_at','reason')})
                if item['retained_fields']:
                    self.issue('review','warning','FIELDS_RETAINED_FROM_PRIOR_OBSERVATION',**loc)
                body, shared = item['body'], item['shared_body']
                media = raw.get('media_types')
                if media is not None and (not isinstance(media, list) or any(not isinstance(v, str) for v in media)):
                    self.issue('validity', 'error', 'INVALID_MEDIA_TYPES', **loc)
                    media = []
                attachments = raw.get('attachment_count')
                if attachments is not None and (not integer(attachments) or attachments < 0):
                    self.issue('validity', 'error', 'INVALID_ATTACHMENT_COUNT', **loc)
                    attachments = None
                is_shared = raw.get('is_shared')
                if is_shared is not None and not isinstance(is_shared, bool):
                    self.issue('validity', 'error', 'INVALID_SHARED_FLAG', **loc)
                    is_shared = None
                body_known = raw.get('body_known')
                if body_known is not None and not isinstance(body_known,bool):
                    self.issue('validity','error','INVALID_BODY_KNOWN_FLAG',**loc)
                    body_known = None
                item['body_known'] = body_known
                if body and body.strip():
                    kind = '正文'
                elif body_known is False:
                    kind = '待核验'
                    self.issue('review','warning','BODY_UNAVAILABLE',**loc)
                elif is_shared is True or (shared and shared.strip()):
                    kind = '无配文分享'
                elif media:
                    kind = '媒体无正文'
                elif attachments:
                    kind = '附件无正文'
                else:
                    kind = '待核验'
                    self.issue('review', 'warning', 'EMPTY_BODY_UNCLASSIFIED', **loc)
                if raw.get('body_truncated') is True:
                    self.issue('review', 'warning', 'BODY_TRUNCATED', **loc)
                item.update(classification=kind, is_shared=is_shared, attachment_count=attachments, media_types=media)
                item['row_issues'] = list(dict.fromkeys(i['code'] for i in self.issues[start_issue:]))
                normalized[bucket].append(item)
            normalized[bucket].sort(key=lambda r: iso(r['published_at']).timestamp() if iso(r['published_at']) else float('-inf'), reverse=True)
        for key, bucket in [('in_range_count','posts'),('outside_range_count','outside_range_posts'),('unknown_date_count','undated_posts')]:
            if self.meta.get(key) is not None and self.meta[key] != len(normalized[bucket]):
                self.issue('validity', 'error', 'METADATA_COUNT_MISMATCH', field=key)
        if any(r['date_status'] == 'unknown' for rows in normalized.values() for r in rows):
            self.issue('coverage', 'warning', 'UNKNOWN_DATES_PREVENT_FULL_WINDOW_VERIFICATION')
        return normalized

    def reconcile_snapshot(self, records, normalized, legacy=False):
        """Replay completed observations; failed attempts remain raw evidence only.

        Page 0 is an already-merged anchor. A complete terminal may contain valid
        observations even when it cannot establish further cursor traversal.
        """
        expected, legacy_ids, failed_ids = {}, set(), set()
        excluded_attempts, excluded_observations = 0, 0
        fields = ('url', 'published_at', 'published_at_raw', 'body', 'shared_body',
                  'shares', 'reactions', 'comments', 'shares_raw', 'reactions_raw', 'comments_raw',
                  'collected_at', 'body_known', 'body_truncated', 'is_shared', 'media_types',
                  'attachment_count', 'count_observed_at', 'field_observed_at', 'retained_fields')
        object_fields = ('count_observed_at', 'field_observed_at')
        invalid_journal_rows = 0
        for _, record in records:
            safe_page = record.get('page') if integer(record.get('page')) else None
            observations = record.get('posts')
            audit = record.get('audit')
            audit = audit if isinstance(audit, dict) else {}
            if (record.get('status') != 'ok' or record.get('stream_final') is not True
                    or not isinstance(record.get('has_next'), bool) or record.get('errors') or audit.get('errors')):
                excluded_attempts += 1
                if isinstance(observations, list):
                    excluded_observations += len(observations)
                    failed_ids.update(str(p['post_id']) for p in observations if isinstance(p, dict)
                                      and isinstance(p.get('post_id'), (str, int)) and not isinstance(p.get('post_id'), bool)
                                      and str(p['post_id']).isdigit())
                continue
            if not isinstance(observations, list):
                self.issue('coverage', 'error', 'JOURNAL_POSTS_INVALID', page=safe_page)
                invalid_journal_rows += 1
                continue
            for post in observations:
                if not isinstance(post, dict) or not isinstance(post.get('post_id'), (str, int)) or isinstance(post.get('post_id'), bool) or not str(post['post_id']).isdigit():
                    self.issue('coverage', 'error', 'JOURNAL_POST_ID_INVALID', page=safe_page)
                    invalid_journal_rows += 1
                    continue
                pid = str(post['post_id'])
                if legacy:
                    legacy_ids.add(pid)
                previous = expected.get(pid)
                old_time, new_time = iso((previous or {}).get('collected_at')), iso(post.get('collected_at'))
                if previous and old_time and new_time and old_time > new_time:
                    continue
                value = copy.deepcopy(post)
                if record.get('page') != 0:
                    for key in object_fields:
                        value[key] = dict(value[key]) if isinstance(value.get(key), dict) else {}
                    value['retained_fields'] = []
                    if previous:
                        prior_fields = previous.get('field_observed_at')
                        prior_fields = prior_fields if isinstance(prior_fields, dict) else {}
                        def preserve(field, time_key):
                            value[field] = previous.get(field)
                            observed = prior_fields.get(time_key)
                            if observed is None:
                                observed = previous.get('collected_at')
                            value['field_observed_at'][time_key] = observed
                            value['retained_fields'].append({'field':field, 'from_collected_at':observed,
                                                            'reason':'unavailable_in_new_observation'})
                        if not post.get('url') and previous.get('url'):
                            preserve('url', 'url')
                        if not post.get('published_at') and previous.get('published_at'):
                            preserve('published_at', 'publication_time')
                            value['published_at_raw'] = previous.get('published_at_raw')
                        if post.get('body_known') is False and previous.get('body_known') is not False:
                            preserve('body', 'body')
                            value['body_known'] = True
                        prior_counts = previous.get('count_observed_at')
                        prior_counts = prior_counts if isinstance(prior_counts, dict) else {}
                        for key in METRICS:
                            if key in value and value[key] is None and previous.get(key) is not None:
                                value[key] = previous[key]
                                value['count_observed_at'][key] = prior_counts.get(key) if prior_counts.get(key) is not None else previous.get('collected_at')
                expected[pid] = value
        snapshot_ids = set()
        verified, unmatched, mismatches = 0, 0, 0
        for bucket, rows in normalized.items():
            raw_rows = self.data.get(bucket)
            for row in rows:
                pid = row['post_id']
                if pid:
                    snapshot_ids.add(pid)
                loc = {'bucket':bucket, 'index':row['record_index']}
                start_issue = len(self.issues)
                raw = raw_rows[row['record_index'] - 1] if isinstance(raw_rows, list) else {}
                if not isinstance(raw, dict) or not pid:
                    row['journal_status'] = 'invalid_record'
                    continue
                prior = expected.get(pid)
                if prior is None:
                    unmatched += 1
                    failed_only = pid in failed_ids
                    row['journal_status'] = 'failed_observation_only' if failed_only else 'not_journal_verified'
                    self.issue('validity' if failed_only else 'review', 'error' if failed_only else 'warning',
                               'FAILED_ONLY_OBSERVATION_IN_SNAPSHOT' if failed_only else 'SNAPSHOT_ID_NOT_JOURNAL_VERIFIED', **loc)
                else:
                    differences = []
                    for field in fields:
                        actual, observed = raw.get(field), prior.get(field)
                        if field in object_fields:
                            actual, observed = actual or {}, observed or {}
                        elif field == 'retained_fields':
                            actual, observed = actual or [], observed or []
                        if actual != observed:
                            differences.append(field)
                    date = iso(prior.get('published_at'))
                    expected_bucket = ('undated_posts' if not date else
                                       'posts' if self.start and self.end and self.start <= date <= self.end else 'outside_range_posts')
                    if bucket != expected_bucket:
                        differences.append('record_bucket')
                    if differences:
                        mismatches += 1
                        legacy = pid in legacy_ids
                        row['journal_status'] = 'legacy_requires_review' if legacy else 'mismatch'
                        for field in differences:
                            self.issue('review' if legacy else 'validity', 'warning' if legacy else 'error',
                                       'LEGACY_SNAPSHOT_JOURNAL_MISMATCH' if legacy else 'SNAPSHOT_JOURNAL_FIELD_MISMATCH', field=field, **loc)
                    else:
                        verified += 1
                        row['journal_status'] = 'legacy_observed_match' if pid in legacy_ids else 'verified'
                row['row_issues'] = list(dict.fromkeys(row['row_issues'] + [i['code'] for i in self.issues[start_issue:]]))
        missing = set(expected) - snapshot_ids
        for pid in sorted(missing):
            self.issue('validity', 'error', 'JOURNAL_POST_MISSING_FROM_SNAPSHOT', post_id=pid)
        if legacy_ids:
            self.issue('review', 'warning', 'LEGACY_JOURNAL_MERGE_REVIEW_REQUIRED', record_count=len(legacy_ids))
        strict_mismatch = any(i['code'] == 'SNAPSHOT_JOURNAL_FIELD_MISMATCH' for i in self.issues)
        failed_snapshot_ids = snapshot_ids & (failed_ids - set(expected))
        status = ('invalid' if missing or strict_mismatch or failed_snapshot_ids else
                  'not_fully_verified' if unmatched or mismatches or legacy_ids or invalid_journal_rows else 'verified')
        return {'status':status, 'journal_unique_ids':len(expected), 'snapshot_rows_matched':verified,
                'missing_journal_ids':len(missing), 'mismatched_snapshot_rows':mismatches,
                'not_journal_verified_rows':unmatched, 'legacy_journal_ids':len(legacy_ids),
                'excluded_failed_attempts':excluded_attempts, 'excluded_failed_observations':excluded_observations,
                'failed_only_snapshot_ids':len(failed_snapshot_ids),
                'definition':'按日志顺序重放完整成功响应的正文、日期与数字及缺失回退；失败响应只留在原始日志，不作为正式字段来源。分页完整性独立判断。'}

    def pagination(self, normalized):
        initial_issue = len(self.issues)
        checkpoint = None
        try:
            checkpoint = read_json(self.source / 'pagination-checkpoint.json')
        except (OSError, ValueError):
            self.issue('coverage', 'error', 'CHECKPOINT_MISSING_OR_INVALID')
        records = []
        if not isinstance(checkpoint, dict):
            checkpoint = {}
        elif not self.same_scope(checkpoint) or checkpoint.get('schema_version') not in (2, 3):
            self.issue('coverage', 'error', 'CHECKPOINT_SCOPE_OR_VERSION_MISMATCH')
        histories = checkpoint.get('history_files', [])
        if not isinstance(histories, list) or not histories:
            histories = []
            self.issue('coverage', 'error', 'HISTORY_LIST_MISSING')
        seen_files = set()
        for file_index, name in enumerate(histories):
            if not isinstance(name, str) or name.startswith(('\\\\','//')):
                self.issue('coverage', 'error', 'INVALID_HISTORY_REFERENCE', history_index=file_index)
                continue
            path = Path(name)
            path = path if path.is_absolute() else self.source / path
            path = path.resolve()
            if path.suffix.lower() != '.jsonl' or path in seen_files:
                self.issue('coverage', 'error', 'INVALID_OR_REPEATED_HISTORY_REFERENCE', history_index=file_index)
                continue
            seen_files.add(path)
            try:
                contents = path.read_text(encoding='utf-8-sig')
            except OSError:
                self.issue('coverage', 'error', 'HISTORY_FILE_MISSING', history_index=file_index)
                continue
            if contents and not contents.endswith('\n'):
                self.issue('coverage', 'error', 'TRUNCATED_JOURNAL', history_index=file_index)
            for line_index, line in enumerate(contents.splitlines(), 1):
                if not line.strip():
                    continue
                try:
                    record = json.loads(line)
                except ValueError:
                    self.issue('coverage', 'error', 'INVALID_JOURNAL_LINE', history_index=file_index, line_index=line_index)
                    continue
                if not isinstance(record, dict) or not self.same_scope(record):
                    self.issue('coverage', 'error', 'JOURNAL_SCOPE_MISMATCH', history_index=file_index, line_index=line_index)
                    continue
                if record.get('schema_version') != checkpoint.get('schema_version'):
                    self.issue('coverage', 'error', 'JOURNAL_VERSION_MISMATCH', history_index=file_index, line_index=line_index)
                records.append((file_index, record))
        self.reconciliation = self.reconcile_snapshot(records, normalized, legacy=checkpoint.get('schema_version') == 2)
        # A clean empty terminal or narrowly proven incomplete stream may precede
        # a complete advancing response to the SAME page/request. Preserve every
        # raw attempt; never collapse a normal success or another failure type.
        active_records = []
        superseded_terminal_attempts, superseded_incomplete_attempts = 0, 0
        position = 0
        while position < len(records):
            first = records[position][1]
            end_position = position + 1
            while end_position < len(records):
                candidate = records[end_position][1]
                if candidate.get('page') != first.get('page') or candidate.get('request_cursor') != first.get('request_cursor'):
                    break
                end_position += 1
            group = records[position:end_position]
            newest = group[-1][1]
            prior_kinds = [('terminal' if clean_retry_observation(r, 'terminal') else
                            'incomplete' if clean_retry_observation(r, 'incomplete') else None) for _, r in group[:-1]]
            if len(group) > 1 and all(prior_kinds) and clean_retry_observation(newest, 'advancing'):
                terminal_count, incomplete_count = prior_kinds.count('terminal'), prior_kinds.count('incomplete')
                superseded_terminal_attempts += terminal_count
                superseded_incomplete_attempts += incomplete_count
                if terminal_count:
                    self.issue('review','warning','TERMINAL_RETRY_SUPERSEDED',page=newest.get('page'),attempt_count=terminal_count)
                if incomplete_count:
                    self.issue('review','warning','INCOMPLETE_RETRY_SUPERSEDED',page=newest.get('page'),attempt_count=incomplete_count)
                active_records.append(group[-1])
            else:
                active_records.extend(group)
            position = end_position
        good = []
        page_numbers = set()
        for file_index, record in active_records:
            page = record.get('page')
            if not integer(page) or page < 0:
                self.issue('coverage', 'error', 'INVALID_JOURNAL_PAGE')
                continue
            if record.get('status') != 'ok' or record.get('stream_final') is not True or not isinstance(record.get('has_next'), bool):
                self.issue('coverage', 'error', 'UNCOMMITTED_OR_FAILED_PAGE', page=page)
                continue
            if page in page_numbers:
                self.issue('coverage', 'error', 'JOURNAL_PAGE_REPEATED', page=page)
            page_numbers.add(page)
            good.append((file_index, record))
        expected_pages = self.meta.get('feed_pages')
        if not integer(expected_pages) or expected_pages < 0:
            self.issue('coverage', 'error', 'INVALID_EXPECTED_PAGE_COUNT')
            expected_pages = max(page_numbers, default=0)
        if [r['page'] for _, r in good] != list(range(expected_pages + 1)):
            self.issue('coverage', 'error', 'PAGE_SEQUENCE_OR_INITIAL_ANCHOR_MISSING')
        cross_batch_connections = 0
        previous = None
        for file_index, record in good:
            page = record['page']
            if record['has_next'] and (not isinstance(record.get('next_cursor'), str) or not record['next_cursor'] or record['next_cursor'] == record.get('request_cursor')):
                self.issue('coverage', 'error', 'CONTINUATION_NOT_ADVANCED', page=page)
            if previous:
                previous_file, prior = previous
                if prior['has_next'] is not True or record.get('request_cursor') != prior.get('next_cursor') or not record.get('request_cursor'):
                    self.issue('coverage', 'error', 'PAGE_CONNECTION_BROKEN', page=page)
                if file_index != previous_file:
                    cross_batch_connections += 1
            previous = (file_index, record)
        usable = [r for _, r in good if r['has_next'] is True and isinstance(r.get('next_cursor'), str) and r['next_cursor']]
        if usable:
            last = usable[-1]
            if checkpoint.get('pages') != last['page'] or checkpoint.get('next_page') != last['page'] + 1 or checkpoint.get('cursor') != last['next_cursor']:
                self.issue('coverage', 'error', 'CHECKPOINT_DOES_NOT_MATCH_LAST_USABLE_PAGE')
        try:
            audit_rows = read_json(self.source / 'crawl-audit.json')
        except (OSError, ValueError):
            audit_rows = None
        if not isinstance(audit_rows, list):
            self.issue('coverage', 'error', 'CRAWL_AUDIT_MISSING_OR_INVALID')
            audit_rows = []
        audit_pages = {}
        for row in audit_rows:
            if not isinstance(row, dict):
                self.issue('coverage', 'error', 'CRAWL_AUDIT_ROW_INVALID')
            elif row.get('kind') == 'page':
                if not integer(row.get('page')) or row['page'] < 1:
                    self.issue('coverage', 'error', 'CRAWL_AUDIT_PAGE_INVALID')
                    continue
                if row.get('page') in audit_pages:
                    self.issue('coverage', 'error', 'CRAWL_AUDIT_PAGE_REPEATED', page=row.get('page'))
                audit_pages[row.get('page')] = row
        boundary_run, boundary_ids, date_regressions = 0, set(), 0
        previous_earliest = None
        page_size = checkpoint.get('page_size', 3)
        if not integer(page_size) or page_size < 1:
            page_size = 3
            self.issue('coverage', 'error', 'INVALID_PAGE_SIZE')
        for _, record in good:
            page = record['page']
            if page == 0:
                continue
            posts = record.get('posts')
            if not isinstance(posts, list) or any(not isinstance(p, dict) for p in posts):
                self.issue('coverage', 'error', 'JOURNAL_POSTS_INVALID', page=page)
                posts = []
            ids = [str(p.get('post_id')) for p in posts]
            dates = [iso(p.get('published_at')) for p in posts]
            if any(not value.isdigit() for value in ids):
                self.issue('coverage', 'error', 'JOURNAL_POST_ID_INVALID', page=page)
            for post_index, post in enumerate(posts):
                post_url = post.get('url')
                try:
                    url_parts = urlsplit(post_url or '')
                    post_match = re.fullmatch(r'/groups/([^/]+)/(?:posts|permalink)/(\d+)/?',url_parts.path)
                    valid_url = (group_token(post_url) in self.group_aliases and post_match is not None
                                 and post_match.group(2) == ids[post_index] and not url_parts.username and not url_parts.password)
                except (ValueError, TypeError):
                    valid_url = False
                if not valid_url:
                    self.issue('coverage', 'error', 'JOURNAL_POST_URL_INVALID', page=page)
                raw_epoch = post.get('published_at_raw')
                if raw_epoch is not None:
                    try:
                        numeric_epoch = float(raw_epoch) if not isinstance(raw_epoch,bool) else float('nan')
                        consistent = dates[post_index] is not None and math.isfinite(numeric_epoch) and abs(dates[post_index].timestamp()-numeric_epoch)<=0.001
                    except (ValueError, TypeError, OverflowError):
                        consistent = False
                    if not consistent:
                        self.issue('coverage', 'error', 'JOURNAL_EPOCH_ISO_MISMATCH', page=page)
            audit = audit_pages.get(page)
            if (not audit or audit.get('stream_final') is not True or audit.get('errors')
                or audit.get('has_next') != record['has_next'] or audit.get('ids') != ids
                or not isinstance(audit.get('dates'), list) or [iso(v) for v in audit['dates']] != dates
                or audit.get('count') != len(posts)):
                self.issue('coverage', 'error', 'CRAWL_AUDIT_JOURNAL_MISMATCH', page=page)
            complete_old = (len(posts) >= page_size and len(set(ids)) >= page_size and self.start is not None
                            and all(date is not None and date < self.start for date in dates))
            if complete_old:
                boundary_run += 1
                boundary_ids.update(ids)
            else:
                boundary_run, boundary_ids = 0, set()
            valid_dates = [d for d in dates if d]
            if valid_dates:
                earliest = min(valid_dates)
                if previous_earliest and max(valid_dates) > previous_earliest:
                    date_regressions += 1
                previous_earliest = earliest
        if set(audit_pages) != set(range(1, expected_pages + 1)):
            self.issue('coverage', 'error', 'CRAWL_AUDIT_PAGE_SEQUENCE_MISMATCH')
        required = self.meta.get('required_boundary_pages', 5)
        if not integer(required) or required < 1:
            required = 5
            self.issue('coverage', 'error', 'INVALID_BOUNDARY_REQUIREMENT')
        chain_verified = not any(i['area'] == 'coverage' and i['level'] == 'error' for i in self.issues[initial_issue:])
        boundary = chain_verified and boundary_run >= required and len(boundary_ids) >= required * page_size
        unknown_dates = any(i['code'] == 'UNKNOWN_PUBLICATION_DATE' for i in self.issues)
        reason = self.meta.get('coverage_status')
        if not isinstance(reason, str) or reason not in {'date_boundary_reached','date_boundary_needs_review','feed_end','page_limit','running','stopped','http_error','response_error','empty_pages','pagination_incomplete','cursor_not_advanced'}:
            reason = 'other_or_missing'
        if reason in ('date_boundary_reached', 'date_boundary_needs_review') and not boundary:
            self.issue('coverage', 'warning', 'CLAIMED_BOUNDARY_NOT_PROVEN')
        regression_claim = self.meta.get('date_order_regressions', 0)
        regression_claim = regression_claim if integer(regression_claim) else (len(regression_claim) if isinstance(regression_claim, list) else 0)
        if date_regressions or regression_claim or reason == 'date_boundary_needs_review':
            self.issue('review', 'warning', 'NON_STRICT_DATE_ORDER_OBSERVED')
        time_verified = boundary and not unknown_dates and reason in ('date_boundary_reached','date_boundary_needs_review')
        status = ('time_boundary_needs_review' if date_regressions or regression_claim or reason == 'date_boundary_needs_review' else 'time_boundary_verified') if time_verified else ('feed_exhausted_time_window_unverified' if reason == 'feed_end' else 'partial_or_unverified')
        return {'status':status, 'time_boundary_verified':time_verified, 'page_chain_verified':chain_verified,
                'superseded_terminal_attempts':superseded_terminal_attempts,
                'superseded_incomplete_attempts':superseded_incomplete_attempts,
                'cross_batch_connections_checked':cross_batch_connections, 'pages_expected':expected_pages,
                'pages_verified':len([r for _,r in good if r['page'] > 0]), 'boundary_pages_required':required,
                'consecutive_complete_old_pages':boundary_run, 'unique_old_posts_in_boundary_run':len(boundary_ids),
                'date_order_regressions_observed':date_regressions, 'collector_reason':reason,
                'scope_limit':'本次账号可见列表；不保证已删除、权限限制或账号不可见内容。'}

    def baseline_check(self, normalized):
        if not self.baseline:
            return {'provided':False}
        try:
            baseline = read_json(self.baseline)
        except (OSError, ValueError):
            self.issue('validity','error','BASELINE_UNREADABLE')
            return {'provided':True,'verified':False}
        if not isinstance(baseline, dict) or not isinstance(baseline.get('posts'), list):
            self.issue('validity','error','BASELINE_FORMAT_INVALID')
            return {'provided':True,'verified':False}
        if not self.same_scope(baseline.get('metadata')):
            self.issue('validity','error','BASELINE_SCOPE_MISMATCH')
        old = set()
        for item in baseline['posts']:
            value = item.get('post_id', item.get('贴文ID')) if isinstance(item,dict) else None
            if not isinstance(value,str) or not value.isdigit():
                self.issue('validity','error','BASELINE_ID_INVALID')
            else:
                old.add(value)
        current = {r['post_id'] for r in normalized['posts'] if r['post_id']}
        missing = old - current
        if missing:
            self.issue('validity','error','BASELINE_POSTS_MISSING',missing_count=len(missing))
        return {'provided':True,'previous_unique_ids':len(old),'retained_ids':len(old & current),'missing_ids':len(missing),'new_ids':len(current-old)}

    def run(self):
        normalized = self.normalize()
        coverage = self.pagination(normalized)
        baseline = self.baseline_check(normalized)
        errors = [i for i in self.issues if i['area']=='validity' and i['level']=='error']
        needs_review = bool(self.issues) or coverage['status'] != 'time_boundary_verified'
        report = {'schema_version':1,
            'validity':{'status':'invalid' if errors else 'valid','error_count':len(errors)},
            'coverage':coverage,
            'review':{'status':'needed' if needs_review else 'no_automated_flags','human_review_completed':False,'automatic_audit_only':True},
            'summary':{'record_counts':{k:len(v) for k,v in normalized.items()},
                'unique_ids':len({r['post_id'] for rows in normalized.values() for r in rows if r['post_id']}),
                'content_types':dict(Counter(r['classification'] for rows in normalized.values() for r in rows)),
                'issue_counts':dict(Counter(i['code'] for i in self.issues)),'baseline':baseline,
                'journal_reconciliation':self.reconciliation},
            'issues':self.issues}
        safe_meta = {'group_id':self.group_id,'group_url':None,'start':self.start.isoformat() if self.start else None,
                     'end':self.end.isoformat() if self.end else None,'time_definition':'Facebook 页面返回时间；ISO 字符串保留时区，起止均包含。',
                     'metric_definition':'赞为全部心情；每项数字以 count_observed_at 对应观测时间为准。观测时间缺失时不推定为本行采集时间。空白不等于零；缩写保留约数。',
                     'exported_at':datetime.now(timezone.utc).isoformat(),'validity_status':report['validity']['status'],
                     'coverage_status':coverage['status'],'review_status':report['review']['status'],
                     'classification_definition':'只依据正文与附件结构分类，不推断主题、身份或个人属性。'}
        if group_token(self.meta.get('group_url')):
            p=urlsplit(self.meta['group_url']);safe_meta['group_url']=urlunsplit((p.scheme,p.netloc,p.path,'',''))
        return {'metadata':safe_meta,**normalized},report


def column_value(row, key):
    if '.' in key:
        parent,child=key.split('.',1)
        return (row.get(parent) or {}).get(child)
    return row.get(key)


def csv_rows(rows):
    for row in rows:
        values=[]
        for key,_ in COLUMNS:
            value=column_value(row,key)
            if isinstance(value,list):value='；'.join(value)
            if key=='post_id' and value is not None:value="'"+str(value)
            else:value=table_text(value)
            values.append(value)
        yield values


def utf16_parts(value, limit=30000):
    parts, current, units = [], [], 0
    for char in value:
        size=2 if ord(char)>0xffff else 1
        if units+size>limit:parts.append(''.join(current));current=[];units=0
        current.append(char);units+=size
    if current or not parts:parts.append(''.join(current))
    return parts


def write_xlsx(path, data, report):
    if importlib.util.find_spec('openpyxl') is None:
        return {'status':'unavailable','reason':'openpyxl is not installed; JSON/CSV remain available'}
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    long_rows=[];escaped_characters=0
    for bucket in BUCKETS:
        rows=data[bucket]
        if bucket!='posts' and not rows:continue
        sheet=workbook.create_sheet({'posts':'目标贴文','outside_range_posts':'范围外记录','undated_posts':'日期待核实'}[bucket])
        sheet.append([label for _,label in COLUMNS]);sheet.freeze_panes='D2';sheet.auto_filter.ref=f'A1:{get_column_letter(len(COLUMNS))}{len(rows)+1}'
        for row in rows:
            values=[]
            for key,_ in COLUMNS:
                value=column_value(row,key)
                if isinstance(value,list):value='；'.join(value)
                if isinstance(value,str):
                    value,count=ILLEGAL_XML.subn(lambda m:'\\u%04x'%ord(m.group(0)),value);escaped_characters+=count
                    pieces=utf16_parts(value)
                    if len(pieces)>1:
                        for number,piece in enumerate(pieces,1):long_rows.append([bucket,row['record_index'],row['post_id'],key,number,piece])
                        value=f'[长文本共 {len(pieces)} 段，见长文本页；JSON/CSV 保留全文]'
                values.append(value)
            sheet.append(values)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='303A46')
        for index,(key,_) in enumerate(COLUMNS,1):
            sheet.column_dimensions[get_column_letter(index)].width=65 if key in ('body','shared_body') else (36 if key=='url' else 23)
        for row in range(2,len(rows)+2):sheet.row_dimensions[row].height=60
    if long_rows:
        sheet=workbook.create_sheet('长文本');sheet.append(['分组','序号','贴文ID','字段','分段序号','文字'])
        for row in long_rows:sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
        sheet.column_dimensions['F'].width=100;sheet.freeze_panes='A2'
    sheet=workbook.create_sheet('说明')
    for key,value in [('数据有效性',report['validity']['status']),('覆盖状态',report['coverage']['status']),('待复核',report['review']['status']),
                      ('审核性质','自动审计不等于人工审核完成。'),('范围口径',report['coverage']['scope_limit']),
                      ('保真对账',report['summary']['journal_reconciliation']['status']),
                      ('时间口径',data['metadata']['time_definition']),('互动数字',data['metadata']['metric_definition']),
                      ('分类','仅按正文及附件结构分类；主题分类需用户明确启用并另列为推断。'),
                      ('长正文','超过单元格限制的文字在长文本页分片；JSON/CSV保留全文。'),
                      ('文本保护','XLSX 将文字单元格显式设为文本类型，保留公式样式文字及原有单引号。CSV 添加文本保护前缀；JSON保留原文。XML不允许的控制字符在XLSX显示为转义文本。')]:sheet.append([key,value])
    sheet.column_dimensions['A'].width=20;sheet.column_dimensions['B'].width=100
    for row in sheet:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    workbook.save(path)
    return {'status':'written','long_text_segments':len(long_rows),'xml_characters_escaped':escaped_characters}


def export(source, output, baseline=None, xlsx=False):
    source,output=Path(source).resolve(),Path(output).resolve()
    if source==output:raise ValueError('OUTPUT_MUST_DIFFER_FROM_SOURCE')
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        raise ValueError('OUTPUT_MUST_BE_NEW_OR_EMPTY')
    data,report=Auditor(source,baseline).run()
    output.mkdir(parents=True,exist_ok=True)
    for bucket in BUCKETS:
        if bucket=='posts' or data[bucket]:
            with (output/f'{bucket}.csv').open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.writer(stream,quoting=csv.QUOTE_ALL);writer.writerow([label for _,label in COLUMNS]);writer.writerows(csv_rows(data[bucket]))
    report['exports']={'xlsx':write_xlsx(output/'posts.xlsx',data,report) if xlsx else {'status':'not_requested'}}
    (output/'posts.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    (output/'audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    counts=report['summary']['record_counts'];coverage=report['coverage']
    readme=(f'# Facebook 群组贴文导出\n\n目标分组 {counts["posts"]} 条；范围外 {counts["outside_range_posts"]} 条；日期待核实分组 {counts["undated_posts"]} 条。所有可识别记录逐行保留，重复或异常不会静默删去。\n\n'
            f'- 数据有效性：{report["validity"]["status"]}。\n- 覆盖状态：{coverage["status"]}。\n- 待复核：{report["review"]["status"]}。自动审计不代表人工审核已完成。\n'
            f'- 已校验分页 {coverage["pages_verified"]} 页；连续完整旧页 {coverage["consecutive_complete_old_pages"]} 页，要求 {coverage["boundary_pages_required"]} 页。日期回跳不会使正确正文或计数自动失效。\n\n'
            '请先看 audit.json 的三个独立结论及异常代码，再使用数据。主 CSV 仅含原始目标分组；范围外和日期待核实记录分别保存，JSON 保留三个分组。无效分组结构或非标量异常字段会明确记录错误，不会伪造可用值。\n\n'
            '最终记录会与完整成功分页观测按 collector 的缺失回退规则重放对账：检查 ID、正文、被转发正文、时间与分组、数字及各自观测时间。失败响应只保留在原始日志，不能成为正式字段来源；同页同游标恢复必须满足严格完整性证据。漏行或 schema 3 字段不符会使数据有效性失败。仅有时间边界证据不能替代保真审核；无日志对应的种子记录仍保留并标记 not_journal_verified，schema 2 的旧合并差异需独立复核。详情见 summary.journal_reconciliation 和每行分页记录保真状态。\n\n'
            '时间采用 Facebook 页面返回值，ISO 字符串保留时区；时间范围的起止均包含。点赞心情数包含所有回应，三个计数分别保留自身观测时间。缺失观测时间不会用本行采集时间补齐；继承旧数字仍保留旧观测时间。空白不等于 0，K/M/万等仅在没有精确字段时作为约数保留。\n\n'
            '默认仅分类正文、无配文分享、媒体无正文、附件无正文、待核验。主题分类默认不运行；只有用户明确要求并提供分类规则后，才可另列“推断主题”，不得当作源事实。\n\n'
            'JSON 保留原始文字；CSV 对可能触发公式的文字和长数字 ID 添加单引号前缀，需要原样文字时使用 JSON，避免误删源文本本来就有的单引号。XLSX 显式使用文本单元格，保留公式样式文字与原有单引号。可用 --xlsx 选项在已安装 openpyxl 时生成 XLSX；长文字分片存放，不截断原始 JSON/CSV。每次导出必须使用新目录或空目录，防止误用旧文件。\n\n'
            '范围仅限本次账号可见列表，不保证已删除或不可见内容。交付文件不复制登录资料、原始请求、分页游标或工作日志。\n')
    (output/'README_导出说明.md').write_text(readme,encoding='utf-8')
    return report


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',required=True,type=Path);parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--baseline',type=Path)
    formats=parser.add_mutually_exclusive_group()
    formats.add_argument('--xlsx',dest='xlsx',action='store_true')
    formats.add_argument('--no-xlsx',dest='xlsx',action='store_false')
    parser.set_defaults(xlsx=False)
    args=parser.parse_args(argv)
    try:report=export(args.source,args.output,args.baseline,args.xlsx)
    except (OSError,ValueError,TypeError,KeyError,OverflowError) as error:
        # Do not echo input contents, source paths, or raw exceptions.
        print(json.dumps({'ok':False,'code':'AUDIT_EXPORT_FAILED','error_type':type(error).__name__}))
        return 2
    print(json.dumps({'ok':report['validity']['status']=='valid','validity':report['validity']['status'],
                      'coverage':report['coverage']['status'],'review':report['review']['status'],
                      'target_records':report['summary']['record_counts']['posts']},ensure_ascii=False))
    return 0 if report['validity']['status']=='valid' else 2


if __name__=='__main__':
    sys.exit(main())
